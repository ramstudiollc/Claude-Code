#!/usr/bin/env python3
"""Build + QC for the 30-day Facebook content system.

Single source of truth: data/posts-*.yaml (+ data/sources.yaml).
Generates: 03-calendar/{calendar.csv,calendar.xlsx,calendar.md},
           04-reels/week-N.md, 05-image-posts/week-N.md,
           07-qc/qc-report.md, viewer.html
Exit code 1 if any hard QC check fails.

Usage:  python3 tools/build.py [--start 2026-10-12]
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import itertools
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

SLOTS = {"I1": ("09:30", 1), "R1": ("13:30", 2), "I2": ("18:30", 3), "R2": ("21:00", 4)}
PILLARS = {
    "PA": "Practical AI",
    "AU": "Automation & n8n",
    "HF": "Hidden features & tools",
    "WF": "Workflows & productivity",
    "LIT": "AI literacy & safety",
    "NEWS": "Current developments",
    "SVC": "Work with me (n8n services)",
}
DRIVERS = {
    "hook", "curiosity", "usefulness", "relatability", "novelty",
    "problem_solution", "visual_progression", "payoff", "rewatch", "save_share",
}
CTA_TYPES = {"save", "share", "question", "follow", "message", "none"}
SPEC = {
    "reel": "Reel 9:16 · 1080×1920 · H.264 MP4 · 30 fps · AAC 48 kHz · safe box x65–1015 / y270–1250 · "
            "burned-in captions y1050–1250 · cover title inside centre 1080×1350",
    "image": "Feed image 4:5 · 1440×1800 PNG (sRGB) · text ≥90 px from edges · headline ≥88 px, body ≥44 px · ≤75 words on image",
}
BANNED = [
    "unlock", "unlocks", "unlocking", "elevate", "elevates", "leverage", "leveraging", "seamless", "seamlessly",
    "robust", "game-changer", "game changer", "game-changing", "revolutionize", "revolutionise", "revolutionary",
    "supercharge", "delve", "cutting-edge", "innovative", "world-class", "fast-paced world", "harness",
    "unleash", "skyrocket", "next level", "next-level", "mind-blowing", "insane",
]
BAIT = [
    r"\bcomment (?:\"?yes|\"?1\b|below if)", r"\btype (?:\"?yes|1\b)", r"\btag (?:a|3|three|your) friends?\b",
    r"\bshare if\b", r"\blike if\b", r"\breact with\b", r"\bvote (?:with|by)\b",
]
REQUIRED = ["id", "type", "series", "pillar", "format", "topic", "objective", "hook", "main_idea", "value",
            "cta", "cta_type", "visual_concept", "production_notes", "caption", "hashtags"]
WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'’+\-./%×]*")

# Pacing limits (see 02-strategy/strategy.md §8)
MAX_WPS_BEAT = 3.0      # hard
WARN_WPS_BEAT = 2.8
MAX_WPS_AVG = 2.6
MAX_OST_WORDS = 10
MAX_BEAT_SEC = 6.0
FIRST_BEAT_MAX = 3.0


def words(text: str | None) -> list[str]:
    return WORD.findall(text or "")


def load():
    sources = yaml.safe_load((DATA / "sources.yaml").read_text())
    posts = []
    for f in sorted(DATA.glob("posts-*.yaml")):
        chunk = yaml.safe_load(f.read_text()) or []
        for p in chunk:
            p["_file"] = f.name
        posts.extend(chunk)
    return sources, posts


def parse_t(t: str):
    a, b = re.split(r"\s*[-–]\s*", str(t))
    return float(a), float(b)


def enrich(posts, start: dt.date):
    for p in posts:
        m = re.fullmatch(r"D(\d\d)-([RI][12])", p["id"])
        if not m:
            raise SystemExit(f"bad id {p['id']}")
        p["day"] = int(m.group(1))
        p["slot"] = m.group(2)
        p["time"], p["slot_no"] = SLOTS[p["slot"]]
        p["date"] = start + dt.timedelta(days=p["day"] - 1)
        p["spec"] = SPEC[p["type"]]
    posts.sort(key=lambda p: (p["day"], p["slot_no"]))
    for i, p in enumerate(posts, 1):
        p["post_no"] = i
    return posts


def caption_full(p):
    tags = " ".join("#" + h.lstrip("#") for h in p.get("hashtags", []))
    return (p["caption"].rstrip() + ("\n\n" + tags if tags else "")).strip()


def on_image_words(p):
    d = p.get("design", {})
    parts = [d.get("headline"), d.get("subhead"), d.get("footer")]
    parts += d.get("body", []) or []
    return sum(len(words(x)) for x in parts if x)


# ---------------------------------------------------------------- QC
def qc(posts, sources):
    errors, warnings, info = [], [], []
    E = errors.append
    W = warnings.append

    reels = [p for p in posts if p["type"] == "reel"]
    imgs = [p for p in posts if p["type"] == "image"]
    ids = [p["id"] for p in posts]
    if len(reels) != 60: E(f"Reel count {len(reels)} ≠ 60")
    if len(imgs) != 60: E(f"Image count {len(imgs)} ≠ 60")
    if len(posts) != 120: E(f"Total {len(posts)} ≠ 120")
    dup = [k for k, v in Counter(ids).items() if v > 1]
    if dup: E(f"Duplicate ids: {dup}")
    days = defaultdict(list)
    for p in posts:
        days[p["day"]].append(p)
    for d in range(1, 31):
        slots = sorted(p["slot"] for p in days.get(d, []))
        if slots != ["I1", "I2", "R1", "R2"]:
            E(f"Day {d}: slots {slots} (need I1, I2, R1, R2)")
    extra = set(days) - set(range(1, 31))
    if extra: E(f"Days outside 1–30: {sorted(extra)}")

    for p in posts:
        pid = p["id"]
        for k in REQUIRED:
            v = p.get(k)
            if v in (None, "", []):
                E(f"{pid}: missing {k}")
        if p.get("pillar") not in PILLARS: E(f"{pid}: unknown pillar {p.get('pillar')}")
        if p.get("cta_type") not in CTA_TYPES: E(f"{pid}: bad cta_type {p.get('cta_type')}")
        for s in p.get("sources", []) or []:
            if s not in sources: E(f"{pid}: unknown source key {s}")
        if p.get("pillar") == "NEWS" and not p.get("sources"):
            E(f"{pid}: NEWS post without sources")

        # caption checks
        cap = p.get("caption", "")
        first = cap.strip().split("\n")[0]
        if len(first) > 125: E(f"{pid}: caption first line {len(first)} chars > 125")
        if re.search(r"https?://|www\.", cap): E(f"{pid}: link in caption body (use Messenger / first comment)")
        nh = len(p.get("hashtags", []))
        if not 2 <= nh <= 4: E(f"{pid}: {nh} hashtags (need 2–4)")
        if len(caption_full(p)) > 1500: W(f"{pid}: caption {len(caption_full(p))} chars (long)")

        # copy checks across all text
        blob = json.dumps({k: v for k, v in p.items() if not k.startswith("_") and k not in ("date",)},
                          ensure_ascii=False, default=str).lower()
        for b in BANNED if not p.get("allow_banned") else []:  # allow_banned: posts that teach the ban list
            if re.search(r"(?<![a-z])" + re.escape(b) + r"(?![a-z])", blob):
                E(f"{pid}: banned word '{b}'")
        for rx in BAIT:
            if re.search(rx, blob):
                E(f"{pid}: engagement-bait pattern {rx}")

        if p["type"] == "reel":
            ret = p.get("retention", {}) or {}
            bad = set(ret) - DRIVERS
            if bad: E(f"{pid}: unknown retention drivers {bad}")
            good = [k for k, v in ret.items() if k in DRIVERS and v and len(str(v)) >= 25]
            if len(good) < 5: E(f"{pid}: only {len(good)} substantiated retention drivers (need ≥5)")
            beats = p.get("script", []) or []
            if not beats:
                E(f"{pid}: no script"); continue
            t_prev = 0.0
            total_vo = 0
            for i, b in enumerate(beats):
                try:
                    a, z = parse_t(b["t"])
                except Exception:
                    E(f"{pid}: beat {i+1} bad time {b.get('t')}"); continue
                if abs(a - t_prev) > 0.01: E(f"{pid}: beat {i+1} starts {a}, expected {t_prev}")
                dur = z - a
                if dur <= 0: E(f"{pid}: beat {i+1} non-positive duration")
                if dur > MAX_BEAT_SEC: W(f"{pid}: beat {i+1} is {dur:.1f}s (>{MAX_BEAT_SEC}s, needs visual change inside)")
                n = len(words(b.get("vo")))
                total_vo += n
                if dur > 0:
                    wps = n / dur
                    if wps > MAX_WPS_BEAT: E(f"{pid}: beat {i+1} VO {n}w/{dur:.1f}s = {wps:.2f} w/s > {MAX_WPS_BEAT}")
                    elif wps > WARN_WPS_BEAT: W(f"{pid}: beat {i+1} VO {wps:.2f} w/s (tight)")
                ow = len(words(b.get("ost")))
                if ow > MAX_OST_WORDS: E(f"{pid}: beat {i+1} on-screen text {ow} words > {MAX_OST_WORDS}")
                if not b.get("vis"): E(f"{pid}: beat {i+1} missing visual direction")
                t_prev = z
            if abs(t_prev - float(p.get("runtime", 0))) > 0.01:
                E(f"{pid}: script ends {t_prev}s but runtime {p.get('runtime')}s")
            a0, z0 = parse_t(beats[0]["t"])
            if z0 > FIRST_BEAT_MAX: E(f"{pid}: first beat {z0}s > {FIRST_BEAT_MAX}s (hook must land fast)")
            hw = [w.lower().strip("’'") for w in words(p.get("hook"))][:3]
            first_txt = " ".join(words((beats[0].get("vo") or "") + " " + (beats[0].get("ost") or ""))).lower()
            if hw and not all(w in first_txt for w in hw):
                W(f"{pid}: hook words {hw} not all found in first beat")
            rt = float(p.get("runtime", 0))
            avg = total_vo / rt if rt else 0
            p["_vo_words"], p["_wps"] = total_vo, round(avg, 2)
            if avg > MAX_WPS_AVG: E(f"{pid}: average VO pace {avg:.2f} w/s > {MAX_WPS_AVG}")
            if rt > 90: E(f"{pid}: runtime {rt}s > 90")
            elif rt > 60: W(f"{pid}: runtime {rt}s > 60")
        else:
            d = p.get("design", {}) or {}
            if not d.get("headline"): E(f"{pid}: missing headline")
            if not d.get("body"): E(f"{pid}: missing body copy")
            if not d.get("layout"): E(f"{pid}: missing layout")
            hw = len(words(d.get("headline")))
            if hw > 10: E(f"{pid}: headline {hw} words > 10")
            tw = on_image_words(p)
            p["_img_words"] = tw
            if tw > 75: E(f"{pid}: {tw} words on image > 75")
            if len(d.get("body", [])) > 8: W(f"{pid}: {len(d['body'])} body lines (>8)")

    # weekday words in a topic must match the posting day (e.g. "Friday review" posts on a Friday)
    for p in posts:
        for wd in ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"):
            if wd.lower() in p["topic"].lower() and p["date"].strftime("%A") != wd:
                E(f"{p['id']}: topic mentions {wd} but posts on {p['date']:%A}")

    # same-day rules
    for d, ps in days.items():
        r = [p for p in ps if p["type"] == "reel"]
        im = [p for p in ps if p["type"] == "image"]
        if len(r) == 2 and r[0]["pillar"] == r[1]["pillar"]:
            E(f"Day {d}: both Reels are {r[0]['pillar']}")
        if len(im) == 2 and im[0]["pillar"] == im[1]["pillar"]:
            W(f"Day {d}: both images are {im[0]['pillar']}")
        svc = sum(p["pillar"] == "SVC" for p in ps)
        if svc > 2: E(f"Day {d}: {svc} service posts (max 2)")

    # consecutive-days service reels
    svc_reel_days = sorted({p["day"] for p in reels if p["pillar"] == "SVC"})
    run = 1
    for a, b in zip(svc_reel_days, svc_reel_days[1:]):
        run = run + 1 if b == a + 1 else 1
        if run > 2: W(f"Service Reels on 3+ consecutive days ending Day {b}")

    # topic similarity
    stop = set("the a an to of and or for in on with your you is it this that how what why ai".split())
    def toks(s):
        return {w.lower() for w in words(s)} - stop
    for p, q in itertools.combinations(posts, 2):
        a, b = toks(p["topic"]), toks(q["topic"])
        if a and b:
            j = len(a & b) / len(a | b)
            if j >= 0.5:
                W(f"Similar topics ({j:.2f}): {p['id']} '{p['topic']}' ~ {q['id']} '{q['topic']}'")

    # format variety: both Reels in a day must differ; no Reel format on 3+ consecutive days
    by_day_fmt = defaultdict(list)
    for p in reels:
        by_day_fmt[p["day"]].append(p["format"])
    for d, fmts in by_day_fmt.items():
        if len(fmts) == 2 and fmts[0] == fmts[1]:
            E(f"Day {d}: both Reels use format '{fmts[0]}'")
    for f in {x for v in by_day_fmt.values() for x in v}:
        streak = 0
        for d in range(1, 31):
            streak = streak + 1 if f in by_day_fmt.get(d, []) else 0
            if streak == 3:
                E(f"Reel format '{f}' runs 3 days in a row ending Day {d}")
    return errors, warnings


# ---------------------------------------------------------------- renderers
def md_escape(s):
    return str(s or "").replace("|", "\\|").replace("\n", "<br>")


def src_lines(p, sources):
    out = []
    for k in p.get("sources", []) or []:
        s = sources[k]
        vol = " ⚠️ re-check within 48 h of posting" if s.get("volatile") else ""
        out.append(f"[{s['title']}]({s['url']}) — {s['publisher']}, {s['date']}{vol}")
    return out


def render_reel(p, sources):
    L = []
    L.append(f'<a id="{p['id']}"></a>\n')
    L.append(f"### #{p['post_no']:03d} · {p['id']} · Day {p['day']} · {p['date']:%a %d %b} · {p['time']} BST · REEL · {p['runtime']} s")
    L.append(f"**{p['topic']}**  ")
    L.append(f"Series: {p['series']} · Pillar: {PILLARS[p['pillar']]} · Format: {p['format']}")
    L.append("")
    L.append("| | |\n|---|---|")
    for k, lab in [("objective", "Objective"), ("hook", "Hook"), ("main_idea", "Main idea"), ("value", "Key value"),
                   ("cta", "CTA"), ("visual_concept", "Visual concept")]:
        L.append(f"| **{lab}** | {md_escape(p[k])} |")
    if p.get("loop"):
        L.append(f"| **Loop** | {md_escape(p['loop'])} |")
    L.append(f"| **Pace** | {p.get('_vo_words', '?')} spoken words in {p['runtime']} s = {p.get('_wps', '?')} words/s |")
    L.append(f"| **Spec** | {p['spec']} |")
    L.append("")
    L.append("**Why people keep watching (retention drivers)**")
    for k, v in p["retention"].items():
        L.append(f"- `{k}`: {v}")
    L.append("")
    L.append("**Script**")
    L.append("")
    L.append("| Time (s) | Voiceover | On-screen text | Visual / edit |\n|---|---|---|---|")
    for b in p["script"]:
        L.append(f"| {b['t']} | {md_escape(b.get('vo'))} | {md_escape(b.get('ost'))} | {md_escape(b.get('vis'))} |")
    L.append("")
    L.append("**Caption**")
    L.append("")
    L.append("```text\n" + caption_full(p) + "\n```")
    L.append(f"**Production notes:** {p['production_notes']}")
    s = src_lines(p, sources)
    if s:
        L.append("")
        L.append("**Sources:** " + " · ".join(s))
    L.append("\n---\n")
    return "\n".join(L)


def render_image(p, sources):
    d = p["design"]
    L = []
    L.append(f'<a id="{p['id']}"></a>\n')
    L.append(f"### #{p['post_no']:03d} · {p['id']} · Day {p['day']} · {p['date']:%a %d %b} · {p['time']} BST · IMAGE")
    L.append(f"**{p['topic']}**  ")
    L.append(f"Series: {p['series']} · Pillar: {PILLARS[p['pillar']]} · Format: {p['format']}")
    L.append("")
    L.append("| | |\n|---|---|")
    for k, lab in [("objective", "Objective"), ("hook", "Hook"), ("main_idea", "Main idea"), ("value", "Key value"),
                   ("cta", "CTA"), ("visual_concept", "Visual concept")]:
        L.append(f"| **{lab}** | {md_escape(p[k])} |")
    L.append(f"| **Words on image** | {p.get('_img_words', '?')} |")
    L.append(f"| **Spec** | {p['spec']} |")
    L.append("")
    L.append("**On-image copy (use exactly)**")
    L.append("")
    L.append(f"- **Headline:** {d['headline']}")
    if d.get("subhead"):
        L.append(f"- **Subhead:** {d['subhead']}")
    L.append("- **Body:**")
    for b in d["body"]:
        L.append(f"  - {b}")
    if d.get("footer"):
        L.append(f"- **Footer:** {d['footer']}")
    L.append(f"- **Layout:** {d['layout']}")
    L.append("")
    L.append("**Caption**")
    L.append("")
    L.append("```text\n" + caption_full(p) + "\n```")
    L.append(f"**Production notes:** {p['production_notes']}")
    s = src_lines(p, sources)
    if s:
        L.append("")
        L.append("**Sources:** " + " · ".join(s))
    L.append("\n---\n")
    return "\n".join(L)


WEEKS = [(1, 1, 7), (2, 8, 14), (3, 15, 21), (4, 22, 30)]


def write_docs(posts, sources, start):
    for wk, a, b in WEEKS:
        rs = [p for p in posts if p["type"] == "reel" and a <= p["day"] <= b]
        im = [p for p in posts if p["type"] == "image" and a <= p["day"] <= b]
        d0, d1 = start + dt.timedelta(days=a - 1), start + dt.timedelta(days=b - 1)
        head = f"Days {a}–{b} · {d0:%a %d %b} → {d1:%a %d %b %Y} · times in BST (UTC+6)\n\n"
        (ROOT / "04-reels" / f"week-{wk}.md").write_text(
            f"# Reel scripts — Week {wk}\n\n{head}"
            "Spec for every Reel: " + SPEC["reel"] + ".\n\n"
            "Read the table left to right: say the **Voiceover**, show the **On-screen text** (key words only, inside the safe box), "
            "cut to the **Visual**. Burned-in captions of the voiceover sit at y1050–1250 on every Reel.\n\n---\n\n"
            + "\n".join(render_reel(p, sources) for p in rs))
        (ROOT / "05-image-posts" / f"week-{wk}.md").write_text(
            f"# Image posts — Week {wk}\n\n{head}"
            "Spec for every image: " + SPEC["image"] + ". Design system: `06-production/design-system.md`.\n\n---\n\n"
            + "\n".join(render_image(p, sources) for p in im))


CSV_COLS = ["post_no", "day", "date", "weekday", "time_bst", "id", "type", "series", "pillar", "format", "topic",
            "objective", "hook", "main_idea", "key_value", "cta", "cta_type", "visual_concept", "caption",
            "sources", "spec", "runtime_s", "production_notes", "script_or_copy_file", "status"]


def row(p, sources):
    wk = next(w for w, a, b in WEEKS if a <= p["day"] <= b)
    return {
        "post_no": p["post_no"], "day": p["day"], "date": p["date"].isoformat(), "weekday": f"{p['date']:%a}",
        "time_bst": p["time"], "id": p["id"], "type": p["type"], "series": p["series"],
        "pillar": f"{p['pillar']} — {PILLARS[p['pillar']]}", "format": p["format"], "topic": p["topic"],
        "objective": p["objective"], "hook": p["hook"], "main_idea": p["main_idea"], "key_value": p["value"],
        "cta": p["cta"], "cta_type": p["cta_type"], "visual_concept": p["visual_concept"],
        "caption": caption_full(p),
        "sources": " ; ".join(f"{sources[k]['title']} ({sources[k]['publisher']}, {sources[k]['date']}) {sources[k]['url']}"
                              for k in p.get("sources", []) or []),
        "spec": p["spec"], "runtime_s": p.get("runtime", ""), "production_notes": p["production_notes"],
        "script_or_copy_file": f"{'04-reels' if p['type']=='reel' else '05-image-posts'}/week-{wk}.md#{p['id']}",
        "status": "planned",
    }


def write_calendar(posts, sources, start):
    rows = [row(p, sources) for p in posts]
    with open(ROOT / "03-calendar" / "calendar.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_COLS)
        w.writeheader()
        w.writerows(rows)

    # markdown overview
    L = ["# 30-day calendar (overview)", "",
         f"Start: **{start:%a %d %b %Y}** · 4 posts/day (2 Reels + 2 images) · times BST (UTC+6). "
         "Full detail per post: `calendar.csv` / `calendar.xlsx`; scripts in `04-reels/`; image copy in `05-image-posts/`.", "",
         "| Day | Date | 09:30 Image | 13:30 Reel | 18:30 Image | 21:00 Reel |", "|---|---|---|---|---|---|"]
    by_day = defaultdict(dict)
    for p in posts:
        by_day[p["day"]][p["slot"]] = p
    def cell(p):
        if not p:
            return "—"
        return f"{md_escape(p['topic'])} <sub>`{p['pillar']}`{' · '+str(p['runtime'])+'s' if p['type']=='reel' else ''}</sub>"
    for d in range(1, 31):
        s = by_day[d]
        date = start + dt.timedelta(days=d - 1)
        L.append(f"| {d} | {date:%a %d %b} | {cell(s.get('I1'))} | {cell(s.get('R1'))} | {cell(s.get('I2'))} | {cell(s.get('R2'))} |")
    (ROOT / "03-calendar" / "calendar.md").write_text("\n".join(L) + "\n")

    # xlsx
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Calendar"
    ws.append(CSV_COLS)
    for r in rows:
        ws.append([r[c] for c in CSV_COLS])
    widths = {"topic": 40, "hook": 45, "main_idea": 45, "key_value": 50, "cta": 35, "visual_concept": 45,
              "caption": 60, "sources": 50, "spec": 40, "production_notes": 50, "objective": 30, "format": 28,
              "series": 20, "pillar": 26, "script_or_copy_file": 30}
    hdr_fill = PatternFill("solid", fgColor="1F2937")
    for i, c in enumerate(CSV_COLS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
    reel_fill = PatternFill("solid", fgColor="EEF2FF")
    for r in range(2, len(rows) + 2):
        is_reel = ws.cell(row=r, column=CSV_COLS.index("type") + 1).value == "reel"
        for i in range(1, len(CSV_COLS) + 1):
            c = ws.cell(row=r, column=i)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if is_reel:
                c.fill = reel_fill
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Grid")
    ws2.append(["Day", "Date", "09:30 Image", "13:30 Reel", "18:30 Image", "21:00 Reel"])
    for d in range(1, 31):
        s = by_day[d]
        ws2.append([d, (start + dt.timedelta(days=d - 1)).isoformat()] +
                   [f"[{s[k]['pillar']}] {s[k]['topic']}" if k in s else "" for k in ("I1", "R1", "I2", "R2")])
    for i, wdt in enumerate([6, 12, 45, 45, 45, 45], 1):
        ws2.column_dimensions[get_column_letter(i)].width = wdt
    for row_ in ws2.iter_rows(min_row=1):
        for c in row_:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
    wb.save(ROOT / "03-calendar" / "calendar.xlsx")


def write_qc(posts, sources, errors, warnings):
    reels = [p for p in posts if p["type"] == "reel"]
    imgs = [p for p in posts if p["type"] == "image"]
    L = ["# QC report (automated)", "",
         f"Generated by `tools/build.py` on {dt.date.today().isoformat()}. "
         "Manual QC notes are in `07-qc/manual-qc.md`.", "",
         f"**Result: {'PASS' if not errors else 'FAIL'}** · {len(errors)} errors · {len(warnings)} warnings", ""]
    L += ["## Counts", "", "| Check | Result |", "|---|---|",
          f"| Reels | {len(reels)} |", f"| Images | {len(imgs)} |", f"| Total | {len(posts)} |",
          f"| Days covered | {len({p['day'] for p in posts})} (each with 2 Reels + 2 images) |", ""]
    c = Counter((p["pillar"], p["type"]) for p in posts)
    L += ["## Pillar mix", "", "| Pillar | Reels | Images | Total | Share |", "|---|---|---|---|---|"]
    for k, v in PILLARS.items():
        t = c[(k, 'reel')] + c[(k, 'image')]
        L.append(f"| {k} — {v} | {c[(k,'reel')]} | {c[(k,'image')]} | {t} | {t/len(posts):.0%} |")
    L.append("")
    rt = [p["runtime"] for p in reels]
    wps = [p.get("_wps", 0) for p in reels]
    L += ["## Reel pacing", "", "| Metric | Value |", "|---|---|",
          f"| Runtime range | {min(rt)}–{max(rt)} s (median {sorted(rt)[len(rt)//2]} s) |",
          f"| Reels ≤ 45 s / 46–60 s / > 60 s | {sum(r<=45 for r in rt)} / {sum(45<r<=60 for r in rt)} / {sum(r>60 for r in rt)} |",
          f"| Average spoken pace | {sum(wps)/len(wps):.2f} words/s (max allowed {MAX_WPS_AVG}) |",
          f"| Fastest Reel average | {max(wps):.2f} words/s |",
          f"| Reels with a planned loop | {sum(bool(p.get('loop')) for p in reels)} |",
          f"| Retention drivers per Reel | min {min(len(p['retention']) for p in reels)}, avg {sum(len(p['retention']) for p in reels)/len(reels):.1f} |", ""]
    dc = Counter(k for p in reels for k in p["retention"])
    L += ["Retention driver usage: " + ", ".join(f"`{k}` {v}" for k, v in dc.most_common()), ""]
    iw = [p.get("_img_words", 0) for p in imgs]
    L += ["## Image copy", "", f"Words on image: min {min(iw)}, max {max(iw)}, avg {sum(iw)/len(iw):.0f} (limit 75).", ""]
    fc = Counter(p["format"] for p in posts)
    L += ["## Formats", "", "| Format | Count |", "|---|---|"] + [f"| {k} | {v} |" for k, v in fc.most_common()] + [""]
    sc = Counter(p["series"] for p in posts)
    L += ["## Series", "", "| Series | Count |", "|---|---|"] + [f"| {k} | {v} |" for k, v in sc.most_common()] + [""]
    cc = Counter(p["cta_type"] for p in posts)
    L += ["## CTA types", "", "| CTA | Count |", "|---|---|"] + [f"| {k} | {v} |" for k, v in cc.most_common()] + [""]
    vol = sorted({(p["day"], p["id"], k) for p in posts for k in p.get("sources", []) or [] if sources[k].get("volatile")})
    L += ["## Pre-flight list (volatile facts, re-check ≤ 48 h before posting)", "", "| Day | Post | Source |", "|---|---|---|"]
    L += [f"| {d} | {i} | [{sources[k]['title']}]({sources[k]['url']}) |" for d, i, k in vol] + [""]
    L += ["## Hard checks run", "",
          "Counts (60/60/120) · 2 Reels + 2 images every day · unique IDs · required fields · ≥5 substantiated retention drivers per Reel · "
          f"contiguous script timing that matches runtime · hook lands in first ≤{FIRST_BEAT_MAX:.0f} s · spoken pace ≤{MAX_WPS_BEAT} w/s per beat and ≤{MAX_WPS_AVG} average · "
          f"on-screen text ≤{MAX_OST_WORDS} words per beat · headline ≤10 words · ≤75 words per image · caption first line ≤125 chars · "
          "2–4 hashtags · no links in captions · banned AI-cliché words · engagement-bait patterns · source keys exist · NEWS posts sourced · "
          "two Reels per day from different pillars and different formats · weekday words in topics match the posting day · no Reel format on 3+ consecutive days · ≤2 service posts per day.", ""]
    L += ["## Errors", ""] + ([f"- ❌ {e}" for e in errors] or ["None."]) + [""]
    L += ["## Warnings (reviewed manually — see manual-qc.md)", ""] + ([f"- ⚠️ {w}" for w in warnings] or ["None."]) + [""]
    (ROOT / "07-qc" / "qc-report.md").write_text("\n".join(L))


def write_viewer(posts, sources):
    data = []
    for p in posts:
        q = {k: v for k, v in p.items() if not k.startswith("_")}
        q["date"] = p["date"].isoformat()
        q["caption_full"] = caption_full(p)
        q["pillar_name"] = PILLARS[p["pillar"]]
        q["src"] = [{"t": sources[k]["title"], "u": sources[k]["url"], "p": sources[k]["publisher"],
                     "d": str(sources[k]["date"]), "v": bool(sources[k].get("volatile"))} for k in p.get("sources", []) or []]
        q["wps"] = p.get("_wps")
        data.append(q)
    tpl = (ROOT / "tools" / "viewer_template.html").read_text()
    out = tpl.replace("__DATA__", json.dumps(data, ensure_ascii=False, default=str).replace("</", "<\\/"))
    (ROOT / "viewer.html").write_text(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--start", default="2026-10-12")
    args = ap.parse_args()
    start = dt.date.fromisoformat(args.start)
    sources, posts = load()
    posts = enrich(posts, start)
    errors, warnings = qc(posts, sources)
    write_calendar(posts, sources, start)
    write_docs(posts, sources, start)
    write_qc(posts, sources, errors, warnings)
    if (ROOT / "tools" / "viewer_template.html").exists():
        write_viewer(posts, sources)
    print(f"posts={len(posts)} reels={sum(p['type']=='reel' for p in posts)} images={sum(p['type']=='image' for p in posts)}")
    print(f"errors={len(errors)} warnings={len(warnings)}")
    for e in errors:
        print("ERR ", e)
    for w in warnings:
        print("WARN", w)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
